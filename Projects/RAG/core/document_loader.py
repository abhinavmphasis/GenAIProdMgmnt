import io
import json
import os
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional


@dataclass
class Document:
    """Represents an extracted document unit or page."""
    content: str
    metadata: Dict[str, Any] = field(default_factory=dict)

    @property
    def source(self) -> str:
        return self.metadata.get("source", "unknown")

    @property
    def page(self) -> Optional[int]:
        return self.metadata.get("page", None)


def _clean_text(text: str) -> str:
    """Cleans up excess whitespace and null characters from extracted text."""
    if not text:
        return ""
    text = text.replace("\x00", "")
    lines = [line.rstrip() for line in text.splitlines()]
    # Collapse multiple consecutive empty lines to a single empty line
    cleaned_lines = []
    prev_empty = False
    for line in lines:
        if not line:
            if not prev_empty:
                cleaned_lines.append("")
            prev_empty = True
        else:
            cleaned_lines.append(line)
            prev_empty = False
    return "\n".join(cleaned_lines).strip()


def load_pdf(stream: io.BytesIO, filename: str) -> List[Document]:
    """Extract text from PDF streams page by page."""
    try:
        import pypdf
    except ImportError:
        raise ImportError(
            "Missing 'pypdf' package. Please install it using: pip install pypdf"
        )

    documents: List[Document] = []
    reader = pypdf.PdfReader(stream)
    total_pages = len(reader.pages)
    
    for page_idx, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
            cleaned = _clean_text(text)
            if cleaned:
                documents.append(
                    Document(
                        content=cleaned,
                        metadata={
                            "source": filename,
                            "page": page_idx,
                            "total_pages": total_pages,
                            "type": "pdf",
                        },
                    )
                )
        except Exception:
            continue

    if not documents:
        documents.append(
            Document(
                content="[Empty PDF or Scanned Document without selectable text]",
                metadata={"source": filename, "page": 1, "total_pages": total_pages, "type": "pdf"},
            )
        )
    return documents


def load_docx(stream: io.BytesIO, filename: str) -> List[Document]:
    """Extract text and tables from DOCX streams."""
    try:
        import docx
    except ImportError:
        raise ImportError(
            "Missing 'python-docx' package. Please install it using: pip install python-docx"
        )

    doc = docx.Document(stream)
    parts: List[str] = []
    
    # Paragraphs
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text.strip())
            
    # Tables
    for table_idx, table in enumerate(doc.tables, start=1):
        table_data = []
        for row in table.rows:
            row_vals = [cell.text.strip() for cell in row.cells]
            table_data.append(" | ".join(row_vals))
        if table_data:
            parts.append(f"\n[Table {table_idx}]\n" + "\n".join(table_data))
            
    content = _clean_text("\n\n".join(parts))
    return [
        Document(
            content=content if content else "[Empty DOCX document]",
            metadata={"source": filename, "type": "docx"},
        )
    ]


def load_csv(stream: io.BytesIO, filename: str) -> List[Document]:
    """Extract CSV content into markdown or line summaries."""
    try:
        import pandas as pd
        df = pd.read_csv(stream)
        num_rows, num_cols = df.shape
        if num_rows <= 100:
            csv_text = df.to_markdown(index=False)
        else:
            csv_text = (
                f"CSV Dataset Overview: {num_rows} rows, {num_cols} columns.\n"
                f"Columns: {', '.join(df.columns.astype(str))}\n\n"
                f"Data Preview (First 50 rows):\n"
                f"{df.head(50).to_markdown(index=False)}\n\n"
                f"Summary Statistics:\n{df.describe(include='all').to_markdown()}"
            )
    except Exception:
        stream.seek(0)
        csv_text = stream.read().decode("utf-8", errors="replace")

    return [
        Document(
            content=_clean_text(csv_text),
            metadata={"source": filename, "type": "csv"},
        )
    ]


def load_json(stream: io.BytesIO, filename: str) -> List[Document]:
    """Extract JSON content into formatted structured text."""
    raw = stream.read().decode("utf-8", errors="replace")
    try:
        obj = json.loads(raw)
        formatted = json.dumps(obj, indent=2, ensure_ascii=False)
    except Exception:
        formatted = raw
    return [
        Document(
            content=_clean_text(formatted),
            metadata={"source": filename, "type": "json"},
        )
    ]


def load_txt(stream: io.BytesIO, filename: str) -> List[Document]:
    """Extract plain text / markdown content."""
    raw = stream.read()
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        try:
            text = raw.decode("latin-1")
        except Exception:
            text = raw.decode("utf-8", errors="replace")
            
    ext = os.path.splitext(filename)[1].lower().lstrip(".")
    return [
        Document(
            content=_clean_text(text),
            metadata={"source": filename, "type": ext or "txt"},
        )
    ]


def load_document_from_bytes(file_bytes: bytes, filename: str) -> List[Document]:
    """Loads and extracts text from file bytes based on file extension."""
    ext = os.path.splitext(filename)[1].lower()
    stream = io.BytesIO(file_bytes)

    if ext == ".pdf":
        return load_pdf(stream, filename)
    elif ext in [".docx", ".doc"]:
        return load_docx(stream, filename)
    elif ext == ".csv":
        return load_csv(stream, filename)
    elif ext == ".json":
        return load_json(stream, filename)
    elif ext in [".txt", ".md", ".markdown", ".rst", ".log", ".py", ".js", ".html", ".xml", ".yaml", ".yml"]:
        return load_txt(stream, filename)
    else:
        # Generic text fallback
        return load_txt(stream, filename)


def load_document_from_file(file_path: str) -> List[Document]:
    """Loads document directly from a local filepath."""
    filename = os.path.basename(file_path)
    with open(file_path, "rb") as f:
        return load_document_from_bytes(f.read(), filename)
