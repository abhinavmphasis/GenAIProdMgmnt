"""
Unit tests to verify document extraction, chunking, vector indexing, and retrieval.
"""

import os
import sys

# Add project root to sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from core.document_loader import load_document_from_bytes, Document
from core.chunking import chunk_documents, DocumentChunk
from core.vector_store import VectorStore
from core.llm import OpenRouterClient, RAGPipeline


def test_txt_loading_and_chunking():
    sample_text = (
        "Artificial Intelligence (AI) in medicine is rapidly advancing. "
        "Machine learning models can detect pathologies in medical imaging with high precision.\n\n"
        "Retrieval-Augmented Generation (RAG) combines semantic vector search with LLMs. "
        "This allows models to cite specific research papers accurately and avoid hallucinations."
    )
    docs = load_document_from_bytes(sample_text.encode("utf-8"), "medical_ai.txt")
    assert len(docs) == 1
    assert docs[0].metadata["source"] == "medical_ai.txt"

    chunks = chunk_documents(docs, chunk_size=150, chunk_overlap=30)
    assert len(chunks) >= 2
    assert chunks[0].source == "medical_ai.txt"
    print(f"PASS: Text loading and chunking created {len(chunks)} chunks.")


def test_vector_store_retrieval():
    vs = VectorStore()

    doc1_content = "The Eiffel Tower is a wrought-iron lattice tower on the Champ de Mars in Paris, France."
    doc2_content = "The Great Wall of China is a series of fortifications built across the historical northern borders of China."
    doc3_content = "Quantum computing uses quantum bits or qubits to perform computations exponentially faster for certain algorithms."

    docs = [
        Document(content=doc1_content, metadata={"source": "paris_guide.txt"}),
        Document(content=doc2_content, metadata={"source": "china_guide.txt"}),
        Document(content=doc3_content, metadata={"source": "quantum_physics.txt"}),
    ]

    chunks = chunk_documents(docs, chunk_size=500, chunk_overlap=50)
    vs.add_chunks(chunks)

    # Query 1: Paris query
    results_paris = vs.similarity_search("Where is the Eiffel Tower located?", top_k=2)
    assert len(results_paris) > 0
    assert results_paris[0][0].source == "paris_guide.txt"
    assert "Eiffel Tower" in results_paris[0][0].content
    print(f"PASS: Paris query retrieved top chunk from '{results_paris[0][0].source}' with score {results_paris[0][1]:.4f}")

    # Query 2: Quantum query
    results_quantum = vs.similarity_search("What are qubits and quantum algorithms?", top_k=1)
    assert len(results_quantum) > 0
    assert results_quantum[0][0].source == "quantum_physics.txt"
    print(f"PASS: Quantum query retrieved top chunk from '{results_quantum[0][0].source}' with score {results_quantum[0][1]:.4f}")

    # Test stats
    stats = vs.get_statistics()
    assert stats["total_sources"] == 3
    assert stats["total_chunks"] == 3
    print(f"PASS: Vector store statistics verified ({stats['total_chunks']} chunks, {stats['total_sources']} sources).")


def test_rag_pipeline_prompt_building():
    vs = VectorStore()
    docs = [
        Document(
            content="Project Apollo was the third United States human spaceflight program carried out by NASA.",
            metadata={"source": "apollo.txt", "page": 1},
        )
    ]
    chunks = chunk_documents(docs, chunk_size=500, chunk_overlap=50)
    vs.add_chunks(chunks)

    client = OpenRouterClient(api_key="mock_key_for_testing")
    pipeline = RAGPipeline(vector_store=vs, llm_client=client)

    retrieved = pipeline.retrieve_context("What was Project Apollo?", top_k=1)
    assert len(retrieved) == 1

    messages = pipeline.build_messages(
        query="What was Project Apollo?",
        retrieved_items=retrieved,
    )
    assert len(messages) == 2
    assert messages[0]["role"] == "system"
    assert messages[1]["role"] == "user"
    assert "apollo.txt" in messages[1]["content"]
    assert "NASA" in messages[1]["content"]
    print("PASS: RAG Pipeline prompt formatting and context injection verified.")


def test_multiformat_loaders():
    import io
    import pypdf
    import docx

    # 1. Test CSV
    csv_bytes = b"id,product,price,stock\n1,Laptop,1200,15\n2,Keyboard,80,120\n3,Mouse,45,200"
    csv_docs = load_document_from_bytes(csv_bytes, "products.csv")
    assert len(csv_docs) == 1
    assert "Laptop" in csv_docs[0].content
    print("PASS: CSV document parsing verified.")

    # 2. Test JSON
    json_bytes = b'{"name": "NeuralNet", "layers": [64, 32, 1], "activation": "relu"}'
    json_docs = load_document_from_bytes(json_bytes, "model_config.json")
    assert len(json_docs) == 1
    assert "NeuralNet" in json_docs[0].content
    print("PASS: JSON document parsing verified.")

    # 3. Test DOCX
    doc = docx.Document()
    doc.add_heading("Architecture Overview", level=1)
    doc.add_paragraph("The system is composed of microservices communicating over gRPC.")
    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_docs = load_document_from_bytes(docx_io.getvalue(), "architecture.docx")
    assert len(docx_docs) == 1
    assert "microservices" in docx_docs[0].content
    print("PASS: DOCX document parsing verified.")

    # 4. Test PDF
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    pdf_io = io.BytesIO()
    writer.write(pdf_io)
    pdf_docs = load_document_from_bytes(pdf_io.getvalue(), "blank.pdf")
    assert len(pdf_docs) == 1
    assert pdf_docs[0].metadata["type"] == "pdf"
    print("PASS: PDF document handling verified.")


if __name__ == "__main__":
    print("--- Running RAG Unit Tests ---")
    test_txt_loading_and_chunking()
    test_vector_store_retrieval()
    test_rag_pipeline_prompt_building()
    test_multiformat_loaders()
    print("\nAll RAG test assertions passed successfully!")

